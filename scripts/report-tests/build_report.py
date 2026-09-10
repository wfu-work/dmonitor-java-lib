"""Build the Chinese DOCX and Markdown report from preserved measurements."""
from pathlib import Path
import json
import math
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
EVIDENCE=ROOT/'docs/test-evidence/2026-09-10'
S=json.loads((EVIDENCE/'summary.json').read_text())
A=json.loads((EVIDENCE/'data-audit.json').read_text())
F=json.loads((EVIDENCE/'fixture-manifest.json').read_text())
D=Document(); MD=[]
SEC=D.sections[0]
SEC.page_width=Cm(21);SEC.page_height=Cm(29.7)
SEC.top_margin=Cm(1.8);SEC.bottom_margin=Cm(1.65)
SEC.left_margin=Cm(1.9);SEC.right_margin=Cm(1.9)
SEC.header_distance=Cm(.7);SEC.footer_distance=Cm(.7)
FONT='Arial Unicode MS'


def font(run,size=None,bold=None,color='000000'):
    run.font.name=FONT
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT)
    if size:run.font.size=Pt(size)
    if bold is not None:run.bold=bold
    run.font.color.rgb=RGBColor.from_string(color)


for name,size in [('Normal',10.5),('Title',22),('Heading 1',16),('Heading 2',12)]:
    style=D.styles[name];style.font.name=FONT;style.font.size=Pt(size)
    style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT)
    style.paragraph_format.space_after=Pt(7)
    style.paragraph_format.line_spacing=1.13
    if name.startswith('Heading'):
        style.font.bold=True;style.paragraph_format.space_before=Pt(9)
        style.paragraph_format.keep_with_next=True
hp=SEC.header.paragraphs[0];hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
font(hp.add_run('DMonitor Java SDK  |  测试报告 v1.2'),8.5)
fp=SEC.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
font(fp.add_run('2026-09-10  |  '),8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');fp._p.append(fld)


def p(text,small=False):
    x=D.add_paragraph(text)
    for r in x.runs:font(r,9 if small else 10.5)
    MD.append(text+'\n');return x


def h(text,level=2):
    x=D.add_paragraph(text,style=f'Heading {level}')
    for r in x.runs:font(r,bold=True)
    MD.append('#'*(level+1)+' '+text+'\n')


def page(title):
    D.add_page_break();h(title,1)


def table(headers,rows,widths=None,size=9):
    t=D.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=widths or [17.2/len(headers)]*len(headers)
    for col,width in zip(t.columns,widths):col.width=Cm(width)
    for row,values in zip([t.rows[0]]+[t.add_row() for _ in rows],[headers]+rows):
        is_header=row==t.rows[0]
        # Row identity is not stable across python-docx wrappers; derive by index below.
        for cell,value,width in zip(row.cells,values,widths):
            cell.width=Cm(width);cell.text=str(value);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for i,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();keep=OxmlElement('w:cantSplit');trpr.append(keep)
        if i==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for j,cell in enumerate(row.cells):
            pr=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd')
            shade.set(qn('w:fill'),'234E64' if i==0 else ('F2F5F7' if i%2==0 else 'FFFFFF'));pr.append(shade)
            margins=OxmlElement('w:tcMar')
            for side in ('top','left','bottom','right'):
                x=OxmlElement('w:'+side);x.set(qn('w:w'),'75');x.set(qn('w:type'),'dxa');margins.append(x)
            pr.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_before=Pt(1);para.paragraph_format.space_after=Pt(1)
                para.paragraph_format.line_spacing=1.08
                if len(str(cell.text))<17:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in para.runs:font(r,size,i==0,'FFFFFF' if i==0 else '000000')
    borders=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
    t._tbl.tblPr.append(borders)
    MD.extend(['| '+' | '.join(map(str,headers))+' |','| '+' | '.join(['---']*len(headers))+' |'])
    MD.extend('| '+' | '.join(str(v).replace('\n','；') for v in row)+' |' for row in rows);MD.append('')
    after=D.add_paragraph();after.paragraph_format.space_after=Pt(1);after.paragraph_format.space_before=Pt(0)


def fmt(x,n=2):return f'{x:.{n}f}'
def v3(xs):return ' / '.join(fmt(x) for x in xs)
def err(s):return '无' if not s['result']['errMsg'] else '有告警'


title=D.add_paragraph('DMonitor Java SDK\n测试报告',style='Title')
for r in title.runs:font(r,22,True)
MD.append('# DMonitor Java SDK 测试报告\n')
p('报告 v1.2  |  SDK 1.0.0  |  原生库 2.0.6.8  |  2026-09-10')
h('结论摘要')
p('结论为条件通过。本轮补充真实 GNSS 数据解算、定位重复性、位移响应相关观测、5 MB 数据处理以及多任务并发测试。独立监测站和基准站已获得固定解，填补了原报告仅用同站数据验证调用链的缺口；绝对定位精度和真实位移反应时间仍需独立真值验证。')
payload=S['payload'];ep=S['epoch_series'];conc=S['concurrency']
table(['测试内容','本轮结果','结论'],[
    ['现有 JUnit 回归','3 项执行，0 失败，0 错误，0 跳过','通过但缺少业务断言'],
    ['数据完整性','155 份样本，316420 帧；净载荷 CRC 失败 0','抽查通过'],
    ['独立站点解算','8 个小时窗口均有结果；严格质量条件 7/8 达标','部分通过'],
    ['定位精度','固定解 ENU 标准差 2.84 / 1.65 / 5.55 mm','重复性已测，绝对精度待验'],
    ['位移反应时间','15 s 采样；首次固定解间隔 60 s','收敛已观测，位移响应待验'],
    ['5 MB 数据处理',f'双站约 5 MB 中位耗时 {fmt(payload[0]["call_ms"]["median"]/1000)} s','含解码的全流程实测'],
    ['多任务并发','1 / 2 / 4 / 8 线程；与串行基线一致','样本内一致性通过'],
],[3.15,8.2,5.85])
h('适用范围')
p('本报告面向 SDK 集成、测试和验收人员。实测范围为当前 macOS arm64 主机、原生库版本及选定数据集。性能数字用于建立本机基线，尚无经确认的时延、容量或精度验收指标，不能据此签署精度达标或容量承诺。')
p('本次保留历史构建结论及已知风险，并区分历史记录、本轮测量和待执行验收方法。正式判定同时要求结果有效、质量达标和性能达标，不能以 JUnit 绿色状态或回调成功代替业务验收。')

page('一 测试环境与计量方法')
table(['项目','本轮基线'],[
    ['Git 提交',S['environment']['git']],['操作系统',f'macOS {S["environment"]["os"]} arm64'],
    ['硬件','Apple M2，8 个逻辑 CPU，24 GiB 内存'],['Java','Oracle JDK 21.0.9 LTS；编译 release 20'],
    ['Maven','Wrapper 3.9.16；本轮统一用 ./mvnw'],['框架','Spring Boot 3.2.4；JNA 5.12.1；Surefire 3.1.2'],
    ['SDK 与算法','com.navfirst:dmonitor-java-lib:1.0.0；原生版本 2.0.6.8'],
    ['JVM 设置','独立 Java 子进程，-Xmx2g；每档单独启动'],
    ['核心参数','rtMode=3，outMode=1，sample=0，vrs=0，navSys=1,4,8,32，bds=1，minFixedRate=0.75'],
    ['逐历元试验','另用 rtMode=2、outMode=0、唯一 outFile；与静态汇总解分开统计'],
],[3.25,13.95],9)
h('计时与统计口径')
p('使用 System.nanoTime()。callMs 从 startMonitor 调用前至同步返回后，包含参数转换、JNA 内存复制、原生解码及解算、结果解析与回调；文件读取、JVM 启动和原生库加载不计入。服务同步调用一次，不能把结果中的 offTime 当作端到端响应时间。')
p('同数据重复执行 6 次，首次单列，后 5 次统计性能。并发每档 8 个任务执行 3 轮，首轮预热，后 2 轮统计；加锁对照执行 2 轮。5 MB 各场景执行 4 次，首轮预热，后 3 次统计。P95 使用最近秩法 ceil(0.95×n)，样本少时等于最大值。')
p('整批耗时从提交该批至全部完成；吞吐量为任务总数除以整批耗时之和。queueMs+callMs 包含排队；加锁对照的 callMs 还包含锁等待。RSS 由 ps 每约 0.2 s 采样，记录整个 JVM 的峰值，包含堆和原生内存，不等于内存泄漏量。')
p('各档基准按顺序执行。本机未采用专用性能环境或固定 CPU 频率；桌面负载、JIT、GC 和温控均可能影响耗时。小样本 P95 不代表生产尾延迟。')

page('二 数据来源与完整性')
p('数据根目录为 /Users/wfu/Downloads/gnss。目录盘点不等于全量解码验证；本轮先盘点全部文件，再抽取第 249 天 13 时的全部 155 个站点文件做帧级核验。')
table(['数据类别','文件数','规模或范围'],[
    ['RTCM 原始文件',A['raw_files'],f'{A["raw_bytes"]:,} 字节；157 个站点'],
    ['RTCM 日期范围','5 天','2026 年第 247 至 251 天，即 09-04 至 09-08'],
    ['RINEX 广播星历',A['nav_files'],f'{A["nav_bytes"]:,} 字节；RINEX 3.02 混合导航文件'],
    ['帧级抽查','155 份','2026/249/13；原文件合计 34,497,349 字节'],
    ['净 RTCM 核验','316,420 帧','32,597,873 字节；CRC 错误、遗留字节均为 0'],
],[3.7,2.3,11.2])
h('传输封装与星历')
p('原始文件包含“十六进制长度 + CRLF + 数据块 + CRLF”的分块传输封装。本轮校验每个块的长度和边界后移除封装，净 RTCM 再执行 CRC24Q 核验。原文件中的非帧字节属于该封装，不应全部归为 RTCM 损坏；CRC 正确也不代表载波观测无周跳、多路径或定位误差。')
p('抽查消息包括 1005/1006 站坐标、1033 设备信息、GPS/GLONASS/Galileo/BDS 等 MSM4 观测，以及 1019、1020、1042、1044、1045、1046 星历消息。独立站点解算使用 BRDM2490.rnx；跨日大样本将 BRDM2490/2500 的导航记录合并，保留一份文件头。')
h('主测试站点与时间')
table(['用途','站点或文件','选用窗口与字节数'],[
    ['监测站','PSYCMM0010','249/13；218,990 字节原文件'],
    ['基准站','PSYCMM0009','249/13；224,473 字节原文件'],
    ['广播星历','nav/2026/BRDM2490.rnx','2026-09-06 对应导航记录'],
    ['时间核验','GPS MSM4 1074','13:00:30 至 14:00:15；239 个历元，间隔 15 s'],
],[2.45,6.15,8.6])
p('测试窗口 13:00:00 至 14:00:00 内原生报告 238 个历元，窗口外 14:00:15 不计入，历元率为 0.9917。时间按 GPS MSM 和原生输出对齐，不能直接把目录小时解释为北京时间；BDS 时标也不能未经转换直接与 GPS 毫秒字段比较。RTCM 1005/1006 坐标仅作站点识别及接收机参考信息，不作为独立测量真值。',True)

page('三 回归结果与独立站点解算')
h('现有 JUnit 回归')
table(['用例','耗时','本轮观察'],[
    ['testUUid','0.469 s','原生调用返回；机器标识在证据日志中脱敏'],
    ['testVersion','0.001 s','算法版本 2.0.6.8'],
    ['testMonitor','0.487 s','原硬编码同站夹具仍会产生固定率业务告警'],
],[4.0,2.0,11.2])
p('执行 ./mvnw test：3 项执行，Failures=0、Errors=0、Skipped=0，BUILD SUCCESS。该测试类仍无 assert 系列业务断言；本轮新增的是独立采集程序与证据，不增加该 JUnit 计数，也不改变生产算法。原生标准输出仍触发 Surefire Corrupted channel 告警。')
h('八个独立小时窗口')
rows=[]
for s in S['hourly_baseline']:
    r=s['result'];hour=s['id'][-2:]
    rows.append([hour+':00',r['solStatus'],fmt(r['fixedRate']*100)+'%',fmt(r['E'],4),fmt(r['N'],4),fmt(r['U'],4),err(s)])
table(['窗口开始','最终状态','固定率','E m','N m','U m','错误'],rows,[2.2,2.2,2.2,2.9,2.6,2.6,2.5],8.8)
p('每个窗口长 1 小时，均为 PSYCMM0010 相对 PSYCMM0009 的解。严格条件为：恰好一次回调、有限坐标、errMsg 为空、fileStatus/navStatus 为 0、solStatus 为 Fixed、fixedRate 不低于 0.75。满足全部条件的窗口为 7/8；10 时窗口最终为 Float，虽固定率为 94.98%，仍不计为通过。')
h('相同输入重复解算')
p(f'13 时窗口重复 6 次均返回 E=60.6892 m、N=21.5883 m、U=6.4236 m，fixedRate=0.9076，Fixed，错误为空。输出显示精度为 0.1 mm，6 次坐标极差均为 0.0000 m；这证明本样本的确定性，不能解释为零定位误差。后 5 次调用中位耗时为 {fmt(S["repeatability"]["call_ms"]["median"])} ms。')

page('四 定位精度与重复性测试')
h('测量目标与当前证据')
p('绝对定位精度要求将解算结果与独立、同期、同坐标框架的参考值比较。应记录参考基站 ECEF、天线相位中心与高度、ENU 原点和轴向、时间系统及真值不确定度。当前数据目录只有原始观测和导航文件，尚无经确认的测站真值或位移真值，因此绝对误差的 RMSE、偏差与 95% 误差暂不作验收结论。')
p('本轮用动态逐历元输出测量相对重复性。样本为 13:00:30 至 14:00:00 的 238 个历元，其中质量码 Q=1 为 213 个（89.50%）、Q=2 为 17 个、Q=0 为 8 个。按 Q=1 列固定解统计，保留其余 25 个历元用于全量异常分析；不能静默删除非固定结果后宣称整体精度达标。')
table(['统计范围','E 标准差 mm','N 标准差 mm','U 标准差 mm'],[
    ['固定解 213 个']+[fmt(v) for v in ep['fixed']['sample_sd_mm']],
    ['全部 238 个']+[fmt(v) for v in ep['all']['sample_sd_mm']],
],[4.6,4.2,4.2,4.2])
table(['相对固定解样本均值','固定解','全量历元'],[
    ['水平距离 P95',fmt(ep['fixed']['horizontal_p95_mm'])+' mm',fmt(ep['all']['horizontal_p95_mm'])+' mm'],
    ['垂向绝对偏离 P95',fmt(ep['fixed']['vertical_p95_mm'])+' mm',fmt(ep['all']['vertical_p95_mm'])+' mm'],
    ['最大水平偏离',fmt(ep['fixed']['horizontal_max_mm'])+' mm',fmt(ep['all']['horizontal_max_mm']/1000,3)+' m'],
    ['最大垂向偏离',fmt(ep['fixed']['vertical_max_mm'])+' mm',fmt(ep['all']['vertical_max_mm']/1000,3)+' m'],
],[7.2,5,5])
p('固定解均值 ENU 为 '+' / '.join(fmt(x,6) for x in ep['fixed_mean_m'])+' m。标准差使用 n−1 分母；P95 按相对固定解样本均值的偏离计算。以上是散布量，不是相对于独立真值的绝对精度；dposmax/dposavg/dposstd 也不能直接替代真值误差。')
h('正式精度验收方法')
p('获得真值后，对每个匹配历元计算 eE=E−E真、eN=N−N真、eU=U−U真；水平误差为 sqrt(eE²+eN²)，三维误差为 sqrt(eE²+eN²+eU²)。分别报告轴向偏差、RMSE、水平/垂向 P95、最大误差、可用率和固定率，并同时列出全量与固定解子集。按短/中/长基线及不少于 3 个独立时段扩展验证；阈值由项目验收要求确定。')

page('五 位移反应时间测试')
h('区分位移响应和解算收敛')
p('位移反应时间应从真实位移发生时刻 t0 开始，至结果首次可靠识别该位移的时刻 tdet，T=t_det−t0。当前样本没有可核验的 t0 和位移幅值，因此不能从静态文件回放耗时、offTime 或首次固定解时间推断真实位移响应。')
p('本轮逐历元样本从 13:00:30 开始，首次 Q=1 出现在 13:01:30，观测间隔为 60 s。这是冷启动样本的首次固定解收敛时间；文件一次性送入时，同步调用完成后才回调，并非按历史时间实时到达。当前采样为 15 s，也不足以验证秒级或亚秒级的动态响应。')
h('不同解算窗口的实测')
rows=[]
for sample in S['windows']:
    r=sample['result'];rows.append([sample['id'].removeprefix('window-'),fmt(sample['callMs']),r['solStatus'],fmt(r['fixedRate']*100)+'%', '通过' if not r['errMsg'] and r['solStatus']=='Fixed' and r['fixedRate']>=.75 else '质量告警'])
table(['窗口','调用耗时 ms','最终状态','固定率','质量判定'],rows,[2.3,3.7,3.1,3.2,4.9])
p('短窗口没有获得与 1 小时窗口相同的质量保障：1/5 分钟还返回历元告警，15 分钟固定率不足 0.75，30/60 分钟满足当前质量条件。此结果只反映该数据的窗口长度与解算质量关系，不能解释为系统必须等待 30 分钟才能识别位移。')
h('待执行的受控位移试验')
p('采用有时间同步的位移平台、全站仪或可信 GNSS 仿真器，设置水平和垂向阶跃，例如 5/10/20/50 mm；这些是建议测试档位，不是已承诺的验收阈值。每档至少重复 10 次，并覆盖冷启动、已固定、失锁后恢复与无位移对照。')
p('预先约定幅值误差带、告警阈值和连续 K 个有效历元确认规则，分别记录首次越阈时间与稳定进入误差带时间。记录采集、接收、排队、计算、回调/告警各时间戳，统计 P50/P95/最大反应时间、漏检率、误报率及恢复时间；无响应的轮次按失败计，不从分母剔除。')
p('filterPeriod 和 processInterval 在当前 Java 实现中只复制到结果，不传入原生结构体，也不建立调度器。不能把这两个参数作为已经生效的滤波窗或轮询周期来计算响应时延。')

page('六 5 MB 数据解码与全流程性能')
h('样本大小与边界')
p('本报告将 5 MB 定义为 5,000,000 字节，另明确 5 MiB=5,242,880 字节。为保持 CRC 和帧完整性，样本取不超过目标大小的最后一个完整帧，实际字节数如下。输入为去除传输封装后的净 RTCM，不重复拼接同一文件，不补零凑整。')
rows=[]
for target,label in [(2500000,'双站合计约 5 MB'),(5000000,'每站约 5 MB')]:
    fs=[f for f in F if f['target_bytes']==target];by={f['station']:f for f in fs}
    rows.append([label,f'{by["PSYCMM0010"]["bytes"]:,}',f'{by["PSYCMM0009"]["bytes"]:,}',f'{sum(f["bytes"] for f in fs):,}'])
table(['场景','监测站字节','基准站字节','合计字节'],rows,[4.9,4.1,4.1,4.1])
p('样本依站点和时间顺序拼接，来源文件、SHA-256、消息类型、有效帧数保存在 fixture-manifest.json；星历另外加载，不计入表中 RTCM 大小。约 5 MB 总量的解算窗口为 09-06 00:00 至 12:00；每站约 5 MB 的窗口为 09-06 00:00 至 09-07 01:00。')
h('实测性能')
rows=[]
for sample in payload:
    rows.append(['约 5 MB 总量' if sample['name'].endswith('2500000') else '约 10 MB 总量',
        fmt(sample['call_ms']['median']/1000),fmt(sample['call_ms']['p95']/1000),
        fmt(sample['throughput_mb_s'],3),fmt(sample['peak_rss_mib'],1)])
table(['场景','中位耗时 s','P95 s','吞吐 MB/s','峰值 RSS MiB'],rows,[4.4,3.3,2.8,3.2,3.5])
rows=[]
for sample in payload:
    r=sample['samples'][-1]['result']
    rows.append(['约 5 MB 总量' if sample['name'].endswith('2500000') else '约 10 MB 总量',
        r['solStatus'],fmt(r['fixedRate']*100)+'%',str(r['roverObsNum'])+' / '+str(r['baseObsNum']),
        '无' if not r['errMsg'] else '有告警'])
table(['场景','状态','固定率','监测/基准历元','业务错误'],rows,[4.4,2.5,2.8,4.6,2.9])
h('判定边界')
p('上表是包含解码的 SDK 全流程耗时，不是纯 RTCM 解码耗时。当前 MonitorLibrary 仅暴露 startGMonitor/getGVersion/getUuid，没有独立解码入口或阶段计时；要验收“5 MB 纯解码性能”，需原生库提供 decode-only 接口或解码阶段埋点。文件读入和准备耗时另存 JSON。')
p('测试完成并返回可用结果，可证明本样本规模下链路可处理；没有约定吞吐和时延阈值，因此性能只记基线。后续补测精确 5 MiB、尾帧截断、CRC 错误、空星历/错日期、10/50 MB、持续重复处理及内存回落，分别判定异常恢复和容量上限。')

page('七 多任务并发解算')
h('任务隔离与公平比较')
p('四档使用完全相同的 8 个不同小时窗口，各自构建 MonitorTask/JNA 内存；共享默认 MonitorServiceImpl 和原生 INSTANCE。全程只注册一次固定回调，通过 extra=id-round 关联结果，避免并发替换回调；不共用 outFile。首轮预热后统计 16 个任务。互斥对照在相同调用外增加同一把锁，计 8 个测量任务。')
rows=[]
for c in conc:
    label=str(c['n'])+(' 加锁' if 'serialized' in c['name'] else '')
    rows.append([label,fmt(c['batch_ms']['median']/1000,3),fmt(c['throughput_tasks_s'],3),
        fmt(c['call_ms']['p95']/1000,3),fmt(c['response_ms']['p95']/1000,3),fmt(c['peak_rss_mib'],1)])
table(['线程数','批次中位 s','任务/s','调用 P95 s','含排队 P95 s','RSS MiB'],rows,[2.2,3.0,2.6,3.0,3.4,3.0],8.6)
p('1 线程的两个测量批次为 11.247 s 和 7.973 s，波动较大；2/4/8 线程的批次约 7.5 至 7.8 s。增加线程并未出现接近线性的吞吐增长，不能由 Java 线程数推定原生算法已并行执行，也不能从这组小样本给出最佳生产并发数。')
h('结果一致性')
p('包含预热在内，四档各 24 次调用、24 次回调，其中严格质量通过 21 次；加锁对照 16 次调用、16 次回调，严格质量通过 14 次。未捕获 Java 异常或子进程崩溃。ENU、固定率、解状态、文件/星历状态、错误、时间、双站历元数及历元率共 13 个字段与串行基线一致。10 时窗口最终 Float 在各档稳定复现，未计为严格质量通过。')
h('仍需验证的并发边界')
p('样本内一致不等于线程安全承诺。默认服务使用未同步的 INSTANCE，虽然另有 synchronizedLibrary 包装 INSTANTCES，默认服务未使用它；MonitorDataServiceImpl 还保存单个可变回调。生产调用应保持回调固定并按任务关联，或在明确原生线程安全前串行保护。后续验证多站点混合负载、异常任务、不同星历/时间范围、8/16/32 档、超时恢复、不同输出文件及至少数小时稳定性。')

page('八 问题清单与验收建议')
table(['编号','级别','问题与影响','关闭条件'],[
    ['TST-001','P1','缺少独立坐标真值，不能验收绝对精度','提供真值、参考框架与阈值，统计全量和固定解误差'],
    ['TST-002','P1','缺少位移 t0/幅值真值，不能验收真实响应时间','完成受控阶跃与无位移对照，记录分阶段时延'],
    ['TST-003','P1','10 时窗口最终 Float；逐历元非固定结果有米级异常偏离','明确质量门禁、异常过滤与告警，并复核观测/解算原因'],
    ['TST-004','P2','没有独立解码计时，不能将总耗时称为解码耗时','提供原生阶段埋点或 decode-only 接口'],
    ['TST-005','P2','默认原生调用未串行保护，回调字段可被覆盖','确认线程安全及回调路由策略，增加压力与异常混跑'],
    ['TST-006','P2','现有 3 个 JUnit 无业务断言且路径硬编码','夹具外置；断言回调、状态、固定率、错误与有限坐标'],
    ['TST-007','P2','无绝对精度/性能 SLO、覆盖率或长期资源基线','确认验收指标，补充 JaCoCo、长时和 Linux 矩阵'],
    ['TST-008','P3','原生输出污染 Surefire 协议，含机器标识','调整日志通道与脱敏；消除 Corrupted channel 告警'],
],[2.3,1.5,7.0,6.4],8.7)
p('P1 表示阻止相关完整业务验收的缺口，不表示所有 SDK 功能不可用；P2 表示影响可靠性、可复现性或回归防护；P3 表示非阻断质量问题。v1.2 按新增验收范围重新整理编号，历史问题说明见附录。',True)
h('发布与交付建议')
p('可确认：当前 macOS arm64 环境能够加载并调用原生库；独立站点和匹配星历可获得固定解；所测 5 MB/10 MB 输入可完成处理；本轮并发样本结果一致。对于精度、真实位移响应、持续运行容量和 Linux 运行时，维持待验状态。')
p('正式门禁至少要求：必需夹具缺失不得被当作通过；无空回调、解析异常或非有限坐标；业务错误为空且文件/星历状态正常；固定率达到约定阈值且最终状态符合场景；精度、响应及性能分别与经确认的指标比较。只记录日志的采集程序不能代替 CI 断言。')
p('推荐先关闭真值和位移试验条件，再确认性能指标和并发部署方式。短基线单站对的毫米级散布不能外推至全部 157 个站点、长基线、复杂遮挡环境或其他硬件平台。')

page('附录 A 复现方法与证据索引')
h('执行步骤')
for text in [
    '1. 在仓库根目录创建 target/report-work；运行 ./mvnw test 保存现有回归结果。',
    '2. 运行 ./mvnw -q -DskipTests test-compile dependency:build-classpath -Dmdep.outputFile=target/report-work/classpath.txt。',
    '3. 运行 scripts/report-tests/rtcm_audit.py，参数依次为 GNSS 数据根目录和 data-audit.json 输出路径。',
    '4. 运行 scripts/report-tests/run_benchmarks.py --data <gnss目录> --license <许可证路径>。需要 JDK 20+ 和有效本机许可证。',
    '5. 运行 scripts/report-tests/summarize.py 计算统计；运行 build_report.py 生成本报告 DOCX 与 Markdown，再渲染并核对 PDF。',
]:p(text)
p('完整可复制命令和计时说明见 scripts/report-tests/README.md。原始 RTCM、星历与许可证不复制到仓库；fixtures 在 target/report-work 中按原文件生成。再次运行会覆盖同名结果，历史证据应先归档。')
h('证据文件')
p('以下文件均位于 docs/test-evidence/2026-09-10，报告数字直接来自这些记录。')
table(['文件','用途'],[
    ['data-audit.json','全目录盘点和 155 份样本的 CRC、消息、时间、参考坐标'],
    ['fixture-manifest.json','大样本字节数、帧边界、源文件列表与 SHA-256'],
    ['repeatability.json','相同输入 6 次调用的完整结果'],
    ['concurrent-1/2/4/8.json','各档任务、计时、回调与业务字段'],
    ['concurrent-8-serialized.json','互斥保护对照'],
    ['window-response.json','1/5/15/30/60 分钟窗口'],
    ['epoch-series.json 与 .pos','逐历元配置、SDK 汇总和 238 条原始位置输出'],
    ['payload-2500000/5000000.json','每站约 2.5 MB / 5 MB 的全流程结果'],
    ['run-manifest.json 与 summary.json','子进程状态、RSS、统计口径和聚合结果'],
    ['maven-test-redacted.log 与各场景 .log','标准回归与实测日志，机器标识已脱敏'],
],[7.1,10.1],8.7)
p('本轮 macOS 原生库 SHA-256：'+S['environment']['native_sha256'],True)

page('附录 B 历史结果与覆盖边界')
h('原 v1.1 报告保留记录')
p('2026-07-29，Git 3b8136f，macOS 26.5.2 arm64，JDK 21.0.9，原生库 2.0.6.8。原报告记录 ./mvnw clean test 执行 3 项、0 失败、0 错误、0 跳过，Spring 上下文 1.580 s，测试总耗时 2.927 s；testMonitor 为 0.589 s。以上属于历史记录，不能替代本轮时间测量。')
p('历史 BRDC、rover 和 base 均使用 GSSK01.2026210binRTCM3，每路 250,039 字节。原回调为 Float、fixedRate=0.0、errMsg=fix-rate less 0.75，双站历元数 239、历元率 0.9958，fileStatus/navStatus 为 0。因此 v1.1 判定为条件通过。用 RTCM 文件充当 RINEX 星历且两站输入相同，不能代表独立站点业务。')
h('历史构建制品')
p('原报告记录 ./mvnw -DskipTests package 成功，18 个主 Java 源文件和 1 个测试源文件完成编译，产物为约 21 MB 的 Spring Boot 可执行 JAR，Main-Class 为 JarLauncher。旧制品 SHA-256 为 1f628672160902383defee525e52869d8d4cb93899041fb494d74000fccc1dc9。该哈希仅标识历史产物，本轮未以该值声称新制品一致。')
table(['平台','资源','覆盖状态'],[
    ['macOS arm64','darwin-aarch64/libDMonitor.dylib','本轮加载和解算已实测'],
    ['Linux arm64','linux-aarch64/libDMonitor.so','历史文件类型与入包检查；本轮未运行'],
    ['Linux x86-64','linux-x86-64/libDMonitor.so','历史文件类型与入包检查；本轮未运行'],
],[3.3,7.3,6.6])
p('历史 macOS 动态依赖包括 CoreFoundation.framework、libSystem.B.dylib、libresolv.9.dylib。Linux ELF 文件存在和资源入包不构成 Linux 运行时兼容性通过依据。')
h('本轮扩展与后续覆盖')
p('本轮补齐独立站点、正确星历、逐历元散布、窗口长度、净 RTCM 大样本、并发一致性及可复现证据。原有回调/参数校验单元测试、时间格式边界、navSys/BDS 转换、广播星历终止符、结果解析异常、覆盖率门槛和 native memory 长时生命周期仍需专项测试。')
p('报告没有执行真实位移平台试验、绝对精度验收、纯解码阶段计时、全天候压力测试或 Linux 算法运行。对这些项目给出的步骤和完成标准为待执行测试方案，不记为已通过。')

D.core_properties.title='DMonitor Java SDK 测试报告'
D.core_properties.subject='定位精度 位移反应时间 5 MB 数据处理 多任务并发'
D.core_properties.version='1.2'
for root in (D.element,D.styles.element):
    for border in root.xpath('.//w:pBdr'):
        border.getparent().remove(border)
D.save(ROOT/'docs/DMonitor-Java-SDK测试报告.docx')
(ROOT/'docs/DMonitor-Java-SDK测试报告.md').write_text('\n'.join(MD))
print('Wrote DOCX and Markdown')
